#!/usr/bin/env python3
"""Genere correspondance_variables.xlsx : le nom de chaque variable exportee.

Le JOB 002/26 n'a pas encore de table « Variable de l'ancienne BD / Variable de
la BD actuelle » comme le JOB 001/26. Seuls les champs d'identification portent
un nom que cette table donne noir sur blanc (NOM_ENQ, DATE_INTERVIEW,
OPERATEUR, SCENARIO...) ; tout le reste a ete ecrit sur ses conventions et reste
a valider.

Ce classeur reprend les variables dans l'ordre du formulaire et dit, pour
chacune, d'ou vient son nom : « table MDS JOB 001/26 » quand il en est repris,
« deduit » -- surligne en jaune -- quand il a ete propose. Seules ces dernieres
lignes sont a relire : corriger la colonne A et renvoyer le fichier suffit, les
valeurs saisies alimentent NOMS_DEDUITS dans build_form.py.

Sortie : correspondance_variables.xlsx
"""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

BASE = Path(__file__).resolve().parent
STRUCTURE = json.loads((BASE / "form_structure.json").read_text(encoding="utf-8"))
SORTIE = BASE / "correspondance_variables.xlsx"

NOMS_BD = STRUCTURE["noms_bd"]
# Noms repris de la table MDS du JOB 001/26 ; les autres ont ete deduits.
NOMS_MDS = set(STRUCTURE["noms_mds"])
# Questions notees : la colonne « Notee » dit lesquelles entrent dans le score.
NOTEES = {n for s in STRUCTURE["notation"].values() for n, _ in s["questions"]}
SECTION_DE = {n: cle for cle, s in STRUCTURE["notation"].items()
              for n, _ in s["questions"]}

A_VALIDER = PatternFill("solid", fgColor="FFF2CC")

wb = Workbook()
ws = wb.active
ws.title = "correspondance"
ws.append(["Nom de variable", "Source", "Libellé / question", "Type", "Notée",
           "Page"])
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="1F5C8B")
    c.alignment = Alignment(horizontal="center")

page = ""
a_valider = 0
for ligne in STRUCTURE["survey"]:
    type_, nom = ligne["type"], ligne["name"]
    if type_ == "begin_group":
        if ligne["appearance"] == "field-list":
            page = ligne["label"]
        continue
    if type_ == "end_group":
        continue
    if nom not in NOMS_BD:                       # notes et calculs internes
        continue
    confirme = nom in NOMS_MDS
    source = "table MDS JOB 001/26" if confirme else "déduit — à valider"
    notee = f"section {SECTION_DE[nom]}" if nom in NOTEES else ""
    ws.append([NOMS_BD[nom], source, ligne["label"].split("\n")[0][:120], type_,
               notee, page])
    if not confirme:
        for colonne in (1, 2):
            ws.cell(row=ws.max_row, column=colonne).fill = A_VALIDER
        a_valider += 1

for colonne, largeur in zip("ABCDEF", (26, 24, 90, 24, 14, 40)):
    ws.column_dimensions[colonne].width = largeur
ws.freeze_panes = "A2"
wb.save(SORTIE)
print(f"OK : {SORTIE.name}")
print(f"  {ws.max_row - 1} variables, {len(NOMS_MDS)} noms repris de la table "
      f"MDS JOB 001/26, {a_valider} déduits à valider (surlignés en jaune)")
print(f"  dont {len(NOTEES)} questions notées")
